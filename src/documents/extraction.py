from pathlib import Path, PurePosixPath

from .config import MAX_EXTRACTED_CHARS

TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")

TEXT_SUFFIXES = frozenset(
    {
        ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl",
        ".yaml", ".yml", ".xml", ".html", ".htm", ".py", ".js", ".ts", ".sql",
        ".ini", ".cfg", ".toml", ".log",
    }
)

PDF_SUFFIXES = frozenset({".pdf"})
WORD_SUFFIXES = frozenset({".docx"})


class UnsupportedDocument(Exception):
    pass


def suffix_of(filename: str) -> str:
    return PurePosixPath(filename.replace("\\", "/")).suffix.lower()


def extract(source: Path, filename: str) -> str:
    suffix = suffix_of(filename)

    if suffix in PDF_SUFFIXES:
        text = _from_pdf(source)
    elif suffix in WORD_SUFFIXES:
        text = _from_word(source)
    elif suffix in TEXT_SUFFIXES:
        text = _from_text(source)
    else:
        raise UnsupportedDocument(f"No extractor for '{suffix or filename}'")

    return _tidy(text)[:MAX_EXTRACTED_CHARS]


def _from_pdf(source: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(source))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _from_word(source: Path) -> str:
    from docx import Document as WordDocument

    document = WordDocument(str(source))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text for cell in row.cells))

    return "\n".join(blocks)


def _from_text(source: Path) -> str:
    raw = source.read_bytes()

    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")

    for encoding in TEXT_ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue

    return raw.decode("utf-8", errors="replace")


def _tidy(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]

    tidied: list[str] = []
    blanks = 0
    for line in lines:
        if line:
            blanks = 0
            tidied.append(line)
            continue
        blanks += 1
        if blanks < 2:
            tidied.append("")

    return "\n".join(tidied).strip()
